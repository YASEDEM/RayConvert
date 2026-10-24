"""RayConvert: local desktop file converter."""
from __future__ import annotations

import argparse
import importlib
import json
import os
import queue
import io
import zipfile
import subprocess
import sys
import tempfile
import threading
import tkinter as tk
from pathlib import Path
from tkinter import filedialog, messagebox, ttk

BASE = Path(__file__).resolve().parent
sys.path.insert(0, str(BASE / "vendor"))
IMAGE = {"png", "jpg", "jpeg", "webp", "bmp", "tif", "tiff", "gif"}
VIDEO = {"mp4", "mov", "mkv", "avi", "webm"}
AUDIO = {"mp3", "wav", "flac", "ogg", "m4a"}
DOC = {"docx", "txt", "html"}

def category(path: Path) -> str:
    ext = path.suffix.lower().lstrip(".")
    for group, extensions in [("image", IMAGE), ("video", VIDEO), ("audio", AUDIO), ("document", DOC), ("pdf", {"pdf"})]:
        if ext in extensions:
            return group
    raise ValueError(f"Unsupported format: {path.suffix or 'file without an extension'}")

def targets(path: Path) -> list[str]:
    ext = path.suffix.lower().lstrip(".")
    group = category(path)
    choices = {
        "image": ["jpg", "png", "webp", "bmp", "tiff", "gif", "pdf"],
        "pdf": ["png", "jpg", "txt", "docx"],
        "document": ["pdf", "docx", "txt", "html"],
        "video": ["mp4", "mov", "mkv", "avi", "webm", "mp3", "wav", "flac", "ogg", "m4a"],
        "audio": ["mp3", "wav", "flac", "ogg", "m4a"],
    }[group]
    normalized = {"jpeg": "jpg", "tif": "tiff"}.get(ext, ext)
    return [item for item in choices if item != normalized]

def verify_output(path: Path, target: str) -> None:
    """Decode the generated file before it becomes visible as a final output."""
    if not path.is_file() or path.stat().st_size == 0:
        raise RuntimeError("The converter produced an empty or incomplete output.")
    if target in IMAGE:
        from PIL import Image
        expected = {"jpg": "JPEG", "tiff": "TIFF"}.get(target, target.upper())
        with Image.open(path) as image:
            if image.format != expected:
                raise RuntimeError("The output content does not match the selected image format.")
            count = getattr(image, "n_frames", 1)
            for frame in range(count):
                image.seek(frame)
                image.load()
    elif target == "pdf":
        import fitz
        with fitz.open(path) as pdf:
            if len(pdf) < 1:
                raise RuntimeError("The PDF contains no pages.")
            for page in pdf:
                page.get_text()
    elif target == "docx":
        from docx import Document
        Document(path)
    elif target in {"txt", "html"}:
        path.read_text(encoding="utf-8")
    elif target in VIDEO | AUDIO:
        import imageio_ffmpeg
        stream = "0:v:0" if target in VIDEO else "0:a:0"
        args = [imageio_ffmpeg.get_ffmpeg_exe(), "-hide_banner", "-v", "error",
                "-xerror", "-i", str(path), "-map", stream, "-f", "null", "-"]
        run_command(args)
    else:
        raise ValueError(f"Unknown target format: {target}")


def publish_files(generated: list[Path], source: Path, target: str, output_dir: Path) -> list[Path]:
    """Publish checked files without overwriting anything already in the output folder."""
    for item in generated:
        verify_output(item, target)
    published = []
    try:
        for index, item in enumerate(generated, 1):
            stem = f"{source.stem} - sayfa {index:03d}" if len(generated) > 1 else source.stem
            for number in range(1, 10000):
                suffix = "" if number == 1 else f" ({number})"
                destination = output_dir / f"{stem}{suffix}.{target}"
                try:
                    if os.name == "nt":
                        os.rename(item, destination)
                    else:
                        os.link(item, destination)
                    published.append(destination)
                    break
                except FileExistsError:
                    continue
            else:
                raise RuntimeError("No available output filename was found.")
        return published
    except Exception:
        for item in published:
            item.unlink(missing_ok=True)
        raise


def run_command(args: list[str]) -> None:
    flags = subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0
    process = subprocess.run(args, capture_output=True, text=True, errors="replace", timeout=1800, creationflags=flags)
    if process.returncode:
        raise RuntimeError((process.stderr or process.stdout).strip()[-900:] or f"Error code: {process.returncode}")

def convert_image(source: Path, target: str, temp: Path) -> list[Path]:
    from PIL import Image, ImageOps
    out = temp / f"result.{target}"
    with Image.open(source) as original:
        if target == "pdf":
            pages = []
            for i in range(getattr(original, "n_frames", 1)):
                original.seek(i)
                pages.append(ImageOps.exif_transpose(original).convert("RGB").copy())
            pages[0].save(out, save_all=True, append_images=pages[1:], resolution=150)
        else:
            frame = ImageOps.exif_transpose(original).copy()
            if target in {"jpg", "bmp"}:
                if frame.mode in {"RGBA", "LA"} or "transparency" in frame.info:
                    rgba = frame.convert("RGBA")
                    background = Image.new("RGB", rgba.size, "white")
                    background.paste(rgba, mask=rgba.getchannel("A"))
                    frame = background
                else:
                    frame = frame.convert("RGB")
            elif target == "gif":
                frame = frame.convert("RGB").quantize(colors=256)
            elif frame.mode not in {"RGB", "RGBA", "L", "LA"}:
                frame = frame.convert("RGBA")
            frame.save(out, format={"jpg": "JPEG", "tiff": "TIFF"}.get(target, target.upper()))
    return [out]

def convert_pdf(source: Path, target: str, temp: Path) -> list[Path]:
    import fitz
    from docx import Document
    outputs = []
    with fitz.open(source) as pdf:
        if pdf.needs_pass:
            raise ValueError("Password-protected PDFs are not supported.")
        if target in {"png", "jpg"}:
            for number, page in enumerate(pdf, 1):
                out = temp / f"page-{number:03d}.{target}"
                page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False).save(str(out))
                outputs.append(out)
        elif target == "txt":
            text = "\n\n".join(page.get_text(sort=True) for page in pdf)
            if not text.strip():
                raise ValueError("The PDF has no selectable text. OCR is required for scanned PDFs.")
            out = temp / "result.txt"
            out.write_text(text, encoding="utf-8")
            outputs.append(out)
        else:
            from docx.shared import Inches
            doc = Document()
            for index, page in enumerate(pdf):
                if index:
                    doc.add_page_break()
                # Keep a visual copy of every page so charts and scans survive.
                picture = io.BytesIO(page.get_pixmap(matrix=fitz.Matrix(1.2, 1.2), alpha=False).tobytes("png"))
                doc.add_picture(picture, width=Inches(6.2))
                text = page.get_text(sort=True).strip()
                if text:
                    doc.add_paragraph("Editable text:")
                    for line in text.splitlines():
                        doc.add_paragraph(line)
            out = temp / "result.docx"
            doc.save(out)
            outputs.append(out)
    return outputs

def convert_document(source: Path, target: str, temp: Path) -> list[Path]:
    """Convert common text documents using bundled Python libraries."""
    from html import escape
    from html.parser import HTMLParser
    from docx import Document
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen import canvas
    from reportlab.lib.utils import simpleSplit

    class TextParser(HTMLParser):
        def __init__(self):
            super().__init__()
            self.parts = []
        def handle_starttag(self, tag, attrs):
            if tag in {"p", "div", "br", "li", "h1", "h2", "h3", "tr"}:
                self.parts.append("\n")
            if tag in {"td", "th"}:
                self.parts.append("  ")
        def handle_data(self, data):
            self.parts.append(data)

    ext = source.suffix.lower().lstrip(".")
    if ext == "docx":
        document = Document(source)
        lines = []
        media = []
        with zipfile.ZipFile(source) as archive:
            for name in archive.namelist():
                if name.startswith("word/media/") and not name.endswith("/"):
                    media.append((name, archive.read(name)))
        for block in document.iter_inner_content():
            if hasattr(block, "rows"):
                for row in block.rows:
                    lines.append(" | ".join(cell.text for cell in row.cells))
            else:
                lines.append(block.text)
    elif ext == "html":
        media = []
        parser = TextParser()
        parser.feed(source.read_text(encoding="utf-8-sig"))
        lines = "".join(parser.parts).splitlines()
    else:
        media = []
        try:
            text = source.read_text(encoding="utf-8-sig")
        except UnicodeDecodeError:
            text = source.read_text(encoding="cp1254")
        lines = text.splitlines()
    lines = [line.strip() for line in lines]
    out = temp / f"result.{target}"
    if target == "txt":
        out.write_text("\n".join(lines), encoding="utf-8")
    elif target == "html":
        body = "\n".join(f"<p>{escape(line)}</p>" if line else "<br>" for line in lines)
        out.write_text("<!doctype html><html lang=\"tr\"><meta charset=\"utf-8\"><body>\n" + body + "\n</body></html>", encoding="utf-8")
    elif target == "docx":
        document = Document()
        for line in lines:
            document.add_paragraph(line)
        document.save(out)
    elif target == "pdf":
        font_file = next((p for p in [Path(r"C:\Windows\Fonts\arial.ttf"), Path(r"C:\Windows\Fonts\segoeui.ttf")] if p.exists()), None)
        if not font_file:
            raise RuntimeError("No Windows TrueType font was found for PDF creation.")
        pdfmetrics.registerFont(TTFont("RayConvertFont", str(font_file)))
        page = canvas.Canvas(str(out), pagesize=A4)
        page.setTitle(source.stem)
        page.setFont("RayConvertFont", 11)
        width, height = A4
        x, y = 48, height - 55
        for line in lines:
            wrapped = simpleSplit(line or " ", "RayConvertFont", 11, width - 96) or [" "]
            for part in wrapped:
                if y < 55:
                    page.showPage()
                    page.setFont("RayConvertFont", 11)
                    y = height - 55
                page.drawString(x, y, part)
                y -= 16
            y -= 5
        if media:
            from PIL import Image as PILImage
            from reportlab.lib.utils import ImageReader
            for name, raw in media:
                try:
                    picture = PILImage.open(io.BytesIO(raw))
                    picture.load()
                except Exception as exc:
                    raise ValueError(f"A Word image could not be added to the PDF: {name}") from exc
                page.showPage()
                page.setFont("RayConvertFont", 10)
                page.drawString(48, height - 50, f"Additional image: {Path(name).name}")
                box_w, box_h = width - 96, height - 130
                scale = min(box_w / picture.width, box_h / picture.height)
                page.drawImage(ImageReader(picture), 48, height - 75 - picture.height * scale,
                               picture.width * scale, picture.height * scale, preserveAspectRatio=True)
        page.save()
    else:
        raise ValueError(f"Unsupported document target format: {target}")
    if not out.exists() or out.stat().st_size == 0:
        raise RuntimeError("The document output could not be created.")
    return [out]

def convert_media(source: Path, target: str, temp: Path) -> list[Path]:
    import imageio_ffmpeg
    out = temp / f"result.{target}"
    args = [imageio_ffmpeg.get_ffmpeg_exe(), "-hide_banner", "-loglevel", "error", "-y", "-i", str(source)]
    video_settings = {
        "mp4": ["-c:v", "libx264", "-pix_fmt", "yuv420p", "-c:a", "aac", "-movflags", "+faststart"],
        "mov": ["-c:v", "libx264", "-pix_fmt", "yuv420p", "-c:a", "aac"],
        "mkv": ["-c:v", "libx264", "-pix_fmt", "yuv420p", "-c:a", "aac"],
        "avi": ["-c:v", "mpeg4", "-c:a", "libmp3lame"],
        "webm": ["-c:v", "libvpx-vp9", "-b:v", "0", "-crf", "32", "-c:a", "libopus"],
    }
    audio_settings = {
        "mp3": ["-vn", "-c:a", "libmp3lame", "-q:a", "2"],
        "wav": ["-vn", "-c:a", "pcm_s16le"],
        "flac": ["-vn", "-c:a", "flac"],
        "ogg": ["-vn", "-c:a", "libvorbis", "-q:a", "5"],
        "m4a": ["-vn", "-c:a", "aac", "-b:a", "192k"],
    }
    run_command(args + (video_settings[target] if target in VIDEO else audio_settings[target]) + [str(out)])
    if not out.exists() or out.stat().st_size == 0:
        raise RuntimeError("FFmpeg did not create an output file.")
    return [out]

def convert(source: Path, target: str, output_dir: Path) -> list[Path]:
    source = source.expanduser().resolve()
    target = target.lower().lstrip(".")
    if not source.is_file():
        raise FileNotFoundError(f"File not found: {source}")
    if target not in targets(source):
        raise ValueError(f"{source.suffix} → .{target} conversion is not supported.")
    output_dir.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix=".rayconvert-", dir=output_dir) as temporary:
        temp = Path(temporary)
        group = category(source)
        if group == "image":
            generated = convert_image(source, target, temp)
        elif group == "pdf":
            generated = convert_pdf(source, target, temp)
        elif group == "document":
            generated = convert_document(source, target, temp)
        else:
            generated = convert_media(source, target, temp)
        if not generated:
            raise RuntimeError("Conversion produced no output.")
        return publish_files(generated, source, target, output_dir)


class RoundedButton(tk.Canvas):
    """Small custom button with restrained rounded corners and hover states."""

    def __init__(self, parent, text, command, variant="secondary", height=38, width=120):
        super().__init__(parent, width=width, height=height, bg=parent.cget("bg"),
                         bd=0, highlightthickness=0, cursor="hand2", takefocus=1)
        self.caption = text
        self.command = command
        self.variant = variant
        self.button_height = height
        self.enabled = True
        self.hovered = False
        self.bind("<Configure>", lambda _event: self.paint())
        self.bind("<Enter>", self.on_enter)
        self.bind("<Leave>", self.on_leave)
        self.bind("<Button-1>", self.on_click)
        self.bind("<Return>", self.on_click)
        self.bind("<space>", self.on_click)
        self.paint()

    def round_rect(self, x0, y0, x1, y1, radius, color):
        r = min(radius, (x1 - x0) / 2, (y1 - y0) / 2)
        self.create_rectangle(x0 + r, y0, x1 - r, y1, fill=color, outline=color)
        self.create_rectangle(x0, y0 + r, x1, y1 - r, fill=color, outline=color)
        for x, y, start in ((x0, y0, 90), (x1 - 2 * r, y0, 0),
                            (x1 - 2 * r, y1 - 2 * r, 270), (x0, y1 - 2 * r, 180)):
            self.create_arc(x, y, x + 2 * r, y + 2 * r, start=start, extent=90,
                            fill=color, outline=color, style="pieslice")

    def paint(self):
        self.delete("all")
        width = max(self.winfo_width(), 2)
        height = max(self.winfo_height(), self.button_height)
        if not self.enabled:
            fill, foreground, border = "#1b1b1d", "#696a70", "#26262a"
        elif self.variant == "primary":
            fill = "#e9e9e7" if self.hovered else "#f5f5f3"
            foreground, border = "#101010", fill
        elif self.variant == "ghost":
            fill = "#1b1b1d" if self.hovered else "#111113"
            foreground, border = "#dedee0", "#28282c"
        else:
            fill = "#222225" if self.hovered else "#171719"
            foreground, border = "#f3f3f4", "#313135"
        self.round_rect(0, 0, width - 1, height - 1, 9, border)
        if border != fill:
            self.round_rect(1, 1, width - 2, height - 2, 8, fill)
        self.create_text(width / 2, height / 2, text=self.caption, fill=foreground,
                         font=("Segoe UI", 10, "bold"))

    def on_enter(self, _event):
        self.hovered = True
        if self.enabled:
            self.configure(cursor="hand2")
        self.paint()

    def on_leave(self, _event):
        self.hovered = False
        self.paint()

    def on_click(self, _event):
        if self.enabled and self.command:
            self.command()

    def set_enabled(self, enabled):
        self.enabled = bool(enabled)
        self.configure(cursor="hand2" if self.enabled else "arrow")
        self.paint()

    def set_text(self, text):
        self.caption = text
        self.paint()


class RayConvertApp:
    BG = "#050505"
    PANEL = "#101011"
    BORDER = "#252527"
    TEXT = "#f2f2f2"
    MUTED = "#96969b"
    SUBTLE = "#68686f"
    SUCCESS = "#a6d99a"
    ERROR = "#e39797"

    def __init__(self, root: tk.Tk):
        self.root = root
        self.files: list[Path] = []
        self.file_states: dict[Path, tuple[str, str]] = {}
        self.events: queue.Queue = queue.Queue()
        self.running = False
        self.target_var = tk.StringVar()
        self.folder_var = tk.StringVar(value=str(Path.home() / "Downloads" / "RayConvert"))
        self.outputs: list[Path] = []
        root.title("RayConvert")
        root.geometry("1100x800")
        root.minsize(900, 760)
        root.configure(bg=self.BG)
        self.build()
        root.after(100, self.enable_dark_titlebar)
        root.after(100, self.poll)

    def enable_dark_titlebar(self):
        if os.name != "nt":
            return
        try:
            import ctypes
            self.root.update_idletasks()
            user32 = ctypes.windll.user32
            dwm = ctypes.windll.dwmapi
            user32.GetParent.argtypes = [ctypes.c_void_p]
            user32.GetParent.restype = ctypes.c_void_p
            dwm.DwmSetWindowAttribute.argtypes = [
                ctypes.c_void_p, ctypes.c_uint, ctypes.c_void_p, ctypes.c_uint
            ]
            hwnd = user32.GetParent(ctypes.c_void_p(self.root.winfo_id()))
            if not hwnd:
                hwnd = self.root.winfo_id()
            for attribute, color in ((20, 1), (34, 0x242424),
                                     (35, 0x050505), (36, 0xF2F2F2)):
                value = ctypes.c_int(color)
                dwm.DwmSetWindowAttribute(
                    ctypes.c_void_p(hwnd), attribute,
                    ctypes.byref(value), ctypes.sizeof(value)
                )
        except Exception:
            pass

    def text(self, parent, value, size=10, color=None, bold=False):
        return tk.Label(parent, text=value, bg=parent.cget("bg"), fg=color or self.TEXT,
                        font=("Segoe UI", size, "bold" if bold else "normal"), anchor="w")

    def separator(self, parent):
        line = tk.Frame(parent, bg=self.BORDER, height=1)
        line.pack(fill="x")
        return line

    @staticmethod
    def readable_size(size):
        if size < 1024:
            return f"{size} B"
        if size < 1024 * 1024:
            return f"{size / 1024:.1f} KB"
        if size < 1024 * 1024 * 1024:
            return f"{size / (1024 * 1024):.1f} MB"
        return f"{size / (1024 * 1024 * 1024):.1f} GB"

    def build(self):
        shell = tk.Frame(self.root, bg=self.BG, padx=36, pady=27)
        shell.pack(fill="both", expand=True)

        header = tk.Frame(shell, bg=self.BG)
        header.pack(fill="x")
        logo = tk.Canvas(header, width=38, height=38, bg=self.BG, bd=0, highlightthickness=0)
        logo.pack(side="left")
        logo.create_rectangle(1, 1, 37, 37, fill="#f3f3f1", outline="")
        logo.create_text(19, 19, text="R", fill="#080808", font=("Segoe UI", 19, "bold"))
        brand = tk.Frame(header, bg=self.BG)
        brand.pack(side="left", padx=(11, 0))
        self.text(brand, "RayConvert", 14, bold=True).pack(anchor="w")
        self.text(brand, "FILE CONVERTER", 8, self.SUBTLE, True).pack(anchor="w")
        privacy = tk.Frame(header, bg=self.BG)
        privacy.pack(side="right")
        tk.Label(privacy, text="●", bg=self.BG, fg=self.SUCCESS,
                 font=("Segoe UI", 9)).pack(side="left", padx=(0, 6))
        self.text(privacy, "LOCAL & OFFLINE", 9, self.MUTED, True).pack(side="left")
        tk.Frame(shell, bg=self.BORDER, height=1).pack(fill="x", pady=(25, 25))

        self.text(shell, "Convert your files.", 26, bold=True).pack(anchor="w")
        self.text(shell, "Select files, choose a format, and save the result.", 10, self.MUTED).pack(anchor="w", pady=(5, 22))

        self.progress = tk.Canvas(shell, height=3, bg="#1c1c1e", bd=0, highlightthickness=0)
        self.progress.pack(side="bottom", fill="x", pady=(14, 0))
        self.progress_fraction = 0
        self.progress.bind("<Configure>", lambda _event: self.draw_progress())
        footer = tk.Frame(shell, bg=self.BG)
        footer.pack(side="bottom", fill="x", pady=(19, 0))
        self.status_label = self.text(footer, "Ready when you are.", 10, self.MUTED)
        self.status_label.pack(side="left")
        self.open_button = RoundedButton(footer, "Open output folder", self.open_output,
                                         variant="ghost", width=146, height=32)
        self.open_button.pack(side="right")
        self.open_button.pack_forget()

        body = tk.Frame(shell, bg=self.BG)
        body.pack(fill="both", expand=True)
        settings_border = tk.Frame(body, bg=self.BORDER, width=302)
        settings_border.pack(side="right", fill="y", padx=(20, 0))
        settings_border.pack_propagate(False)
        settings = tk.Frame(settings_border, bg=self.PANEL, padx=21, pady=18)
        settings.pack(fill="both", expand=True, padx=1, pady=1)
        controls = tk.Frame(settings, bg=self.PANEL)
        controls.pack(side="bottom", fill="x")
        self.convert_button = RoundedButton(controls, "Convert files", self.start,
                                            variant="primary", width=254, height=48)
        self.convert_button.pack(fill="x")
        self.convert_button.set_enabled(False)
        self.text(controls, "© 2026 YASEDEM",
                  8, self.SUBTLE).pack(anchor="center", pady=(10, 0))

        files_border = tk.Frame(body, bg=self.BORDER)
        files_border.pack(side="left", fill="both", expand=True)
        files_panel = tk.Frame(files_border, bg=self.PANEL)
        files_panel.pack(fill="both", expand=True, padx=1, pady=1)
        file_header = tk.Frame(files_panel, bg=self.PANEL, padx=20, pady=16)
        file_header.pack(fill="x")
        title_group = tk.Frame(file_header, bg=self.PANEL)
        title_group.pack(side="left")
        self.text(title_group, "Files", 13, bold=True).pack(side="left")
        self.count_label = self.text(title_group, "0 added", 9, self.SUBTLE)
        self.count_label.pack(side="left", padx=(10, 0), pady=(2, 0))
        self.add_button = RoundedButton(file_header, "+  Add files", self.add_files, width=105, height=34)
        self.add_button.pack(side="right")
        self.clear_button = RoundedButton(file_header, "Clear", self.clear_files, variant="ghost", width=65, height=34)
        self.clear_button.pack(side="right", padx=(0, 8))
        self.separator(files_panel)

        list_container = tk.Frame(files_panel, bg=self.PANEL)
        list_container.pack(fill="both", expand=True)
        self.list_canvas = tk.Canvas(list_container, bg=self.PANEL, bd=0, highlightthickness=0)
        self.list_canvas.pack(side="left", fill="both", expand=True)
        self.scrollbar = tk.Canvas(list_container, width=8, bg=self.PANEL,
                                   bd=0, highlightthickness=0, cursor="hand2")
        self.scrollbar.pack(side="right", fill="y")
        self.scrollbar.bind("<Button-1>", self.on_scrollbar_click)
        self.scrollbar.bind("<Configure>", lambda _event: self.draw_scrollbar())
        self.scroll_first, self.scroll_last = 0.0, 1.0
        self.list_canvas.configure(yscrollcommand=self.update_scrollbar)
        self.list_inner = tk.Frame(self.list_canvas, bg=self.PANEL)
        self.list_window = self.list_canvas.create_window((0, 0), window=self.list_inner, anchor="nw")
        self.list_inner.bind("<Configure>", lambda _event: self.list_canvas.configure(scrollregion=self.list_canvas.bbox("all")))
        self.list_canvas.bind("<Configure>", lambda event: self.list_canvas.itemconfigure(self.list_window, width=event.width))
        self.list_canvas.bind("<Enter>", lambda _event: self.list_canvas.bind_all("<MouseWheel>", self.on_mousewheel))
        self.list_canvas.bind("<Leave>", lambda _event: self.list_canvas.unbind_all("<MouseWheel>"))

        self.separator(files_panel)
        file_foot = tk.Frame(files_panel, bg=self.PANEL, padx=20, pady=13)
        file_foot.pack(fill="x")
        self.text(file_foot, "Files stay on this device", 9, self.SUBTLE).pack(side="left")
        self.text(file_foot, "No uploads", 9, self.SUBTLE).pack(side="right")

        self.text(settings, "OUTPUT SETTINGS", 9, self.MUTED, True).pack(anchor="w")
        tk.Frame(settings, bg=self.BORDER, height=1).pack(fill="x", pady=(13, 16))
        self.text(settings, "FORMAT", 9, self.SUBTLE, True).pack(anchor="w")
        self.format_button = RoundedButton(settings, "Select format    ▾", self.open_format_menu,
                                           width=254, height=44)
        self.format_button.pack(fill="x", pady=(8, 17))
        self.text(settings, "SAVE TO", 9, self.SUBTLE, True).pack(anchor="w")
        folder_border = tk.Frame(settings, bg="#343438", padx=1, pady=1)
        folder_border.pack(fill="x", pady=(8, 6))
        folder_box = tk.Frame(folder_border, bg="#171719", padx=10, pady=9)
        folder_box.pack(fill="x")
        self.folder_entry = tk.Entry(folder_box, textvariable=self.folder_var, bg="#171719",
                                     fg=self.TEXT, insertbackground=self.TEXT, selectbackground="#4a4a50",
                                     relief="flat", bd=0, font=("Segoe UI", 9))
        self.folder_entry.pack(fill="x")
        self.browse_button = RoundedButton(settings, "Browse folders", self.choose_folder,
                                           variant="ghost", width=254, height=35)
        self.browse_button.pack(fill="x")
        tk.Frame(settings, bg=self.BORDER, height=1).pack(fill="x", pady=(17, 11))
        self.hint = self.text(settings, "Originals remain untouched. Existing files are never replaced.",
                              9, self.MUTED)
        self.hint.configure(wraplength=245, justify="left")
        self.hint.pack(anchor="w", fill="x")

        self.render_files()

    def on_mousewheel(self, event):
        self.list_canvas.yview_scroll(int(-event.delta / 120), "units")

    def update_scrollbar(self, first, last):
        self.scroll_first, self.scroll_last = float(first), float(last)
        self.draw_scrollbar()

    def draw_scrollbar(self):
        self.scrollbar.delete("all")
        if self.scroll_last - self.scroll_first >= 0.999:
            return
        height = self.scrollbar.winfo_height()
        top = self.scroll_first * height
        bottom = max(top + 26, self.scroll_last * height)
        self.scrollbar.create_rectangle(2, top, 7, min(bottom, height),
                                        fill="#3a3a3e", outline="")

    def on_scrollbar_click(self, event):
        height = max(1, self.scrollbar.winfo_height())
        self.list_canvas.yview_moveto(max(0, min(1, event.y / height)))

    def draw_progress(self):
        self.progress.delete("all")
        width = self.progress.winfo_width()
        if width > 0 and self.progress_fraction:
            self.progress.create_rectangle(0, 0, width * self.progress_fraction, 3,
                                           fill="#e8e8e6", outline="")

    def render_files(self):
        for child in self.list_inner.winfo_children():
            child.destroy()
        self.count_label.configure(text=f"{len(self.files)} added")
        if not self.files:
            empty = tk.Frame(self.list_inner, bg=self.PANEL, pady=104)
            empty.pack(fill="both", expand=True)
            self.text(empty, "No files yet", 16, bold=True).pack(anchor="center")
            self.text(empty, "Add files to start a conversion.", 10, self.MUTED).pack(anchor="center", pady=(7, 0))
            self.clear_button.set_enabled(False)
            return
        self.clear_button.set_enabled(not self.running)
        badges = {"image": "IMG", "video": "VID", "audio": "AUD",
                  "document": "DOC", "pdf": "PDF"}
        for path in self.files:
            row = tk.Frame(self.list_inner, bg=self.PANEL, padx=18, pady=14)
            row.pack(fill="x")
            badge = tk.Frame(row, bg="#222225", width=44, height=44)
            badge.pack(side="left")
            badge.pack_propagate(False)
            self.text(badge, badges[category(path)], 8, "#c3c3c8", True).pack(expand=True)
            details = tk.Frame(row, bg=self.PANEL)
            details.pack(side="left", fill="x", expand=True, padx=(13, 8))
            display = path.name if len(path.name) <= 45 else path.name[:42] + "..."
            self.text(details, display, 10, bold=True).pack(anchor="w")
            state, detail = self.file_states.get(path, ("Queued", ""))
            if state == "Failed":
                secondary = detail[:80] + ("..." if len(detail) > 80 else "")
            else:
                try:
                    size = self.readable_size(path.stat().st_size)
                except OSError:
                    size = "Unknown size"
                secondary = f"{size}  ·  {path.parent}"
                if len(secondary) > 75:
                    secondary = secondary[:72] + "..."
            self.text(details, secondary, 9, self.ERROR if state == "Failed" else self.SUBTLE).pack(anchor="w", pady=(4, 0))
            right = tk.Frame(row, bg=self.PANEL)
            right.pack(side="right")
            color = self.SUCCESS if state == "Done" else self.ERROR if state == "Failed" else self.SUBTLE
            self.text(right, state, 9, color).pack(side="left", padx=(0, 12))
            remove = tk.Label(right, text="×", bg=self.PANEL, fg=self.SUBTLE,
                              font=("Segoe UI", 15), cursor="hand2")
            remove.pack(side="right")
            if not self.running:
                remove.bind("<Button-1>", lambda _event, item=path: self.remove_file(item))
                remove.bind("<Enter>", lambda _event, widget=remove: widget.configure(fg=self.TEXT))
                remove.bind("<Leave>", lambda _event, widget=remove: widget.configure(fg=self.SUBTLE))
            tk.Frame(self.list_inner, bg=self.BORDER, height=1).pack(fill="x", padx=18)

    def add_files(self):
        if self.running:
            return
        chosen = filedialog.askopenfilenames(title="Select files to convert")
        skipped = 0
        for name in chosen:
            path = Path(name)
            try:
                category(path)
            except ValueError:
                skipped += 1
                continue
            if path not in self.files:
                self.files.append(path)
                self.file_states[path] = ("Queued", "")
        self.render_files()
        self.refresh_targets()
        if skipped:
            self.status_label.configure(text=f"{skipped} unsupported file(s) skipped.", fg=self.ERROR)
        elif chosen:
            self.status_label.configure(text=f"{len(self.files)} file(s) ready.", fg=self.MUTED)

    def remove_file(self, path):
        if self.running:
            return
        self.files.remove(path)
        self.file_states.pop(path, None)
        self.render_files()
        self.refresh_targets()

    def clear_files(self):
        if self.running:
            return
        self.files.clear()
        self.file_states.clear()
        self.render_files()
        self.refresh_targets()
        self.status_label.configure(text="Ready when you are.", fg=self.MUTED)

    def refresh_targets(self):
        previous = self.target_var.get().lower()
        common = set(targets(self.files[0])) if self.files else set()
        for path in self.files[1:]:
            common.intersection_update(targets(path))
        order = ["pdf", "docx", "txt", "html", "jpg", "png", "webp", "bmp", "tiff",
                 "gif", "mp4", "mov", "mkv", "avi", "webm", "mp3", "wav", "flac", "ogg", "m4a"]
        self.available_targets = [item for item in order if item in common]
        chosen = previous if previous in self.available_targets else (
            self.available_targets[0] if self.available_targets else "")
        self.target_var.set(chosen)
        self.format_button.set_text(chosen.upper() + "    ▾" if chosen else "Select format    ▾")
        self.format_button.set_enabled(bool(self.available_targets) and not self.running)
        self.convert_button.set_enabled(bool(self.available_targets) and not self.running)
        self.refresh_hint()
        if self.files and not self.available_targets:
            self.status_label.configure(text="Selected files have no common output format.", fg=self.ERROR)

    def open_format_menu(self):
        if not self.available_targets or self.running:
            return
        menu = tk.Menu(self.root, tearoff=0, bg="#1b1b1d", fg=self.TEXT,
                       activebackground="#323236", activeforeground=self.TEXT,
                       font=("Segoe UI", 10), bd=1, relief="solid")
        for target in self.available_targets:
            menu.add_command(label=target.upper(), command=lambda item=target: self.select_target(item))
        x = self.format_button.winfo_rootx()
        y = self.format_button.winfo_rooty() + self.format_button.winfo_height() + 3
        menu.tk_popup(x, y)
        menu.grab_release()

    def select_target(self, target):
        self.target_var.set(target)
        self.format_button.set_text(target.upper() + "    ▾")
        self.refresh_hint()

    def refresh_hint(self):
        target = self.target_var.get().lower()
        if target == "docx" and any(category(path) == "pdf" for path in self.files):
            text = "PDF pages are embedded as images. Selectable text is added when available."
        elif target in {"pdf", "txt", "html", "docx"} and any(category(path) == "document" for path in self.files):
            text = "Complex document layouts may change. Your original files stay untouched."
        elif target in {"jpg", "gif"}:
            text = "This format can reduce color detail or transparency. Originals stay untouched."
        elif target in VIDEO | AUDIO:
            text = "Media is re-encoded, so quality may change. Originals stay untouched."
        else:
            text = "Originals remain untouched. Existing files are never replaced."
        self.hint.configure(text=text)

    def choose_folder(self):
        if self.running:
            return
        path = filedialog.askdirectory(title="Choose an output folder", initialdir=self.folder_var.get())
        if path:
            self.folder_var.set(path)

    def start(self):
        if self.running or not self.files or not self.target_var.get():
            self.status_label.configure(text="Add files and choose a format first.", fg=self.ERROR)
            return
        folder = self.folder_var.get().strip()
        if not folder:
            self.status_label.configure(text="Choose an output folder.", fg=self.ERROR)
            return
        self.running = True
        self.outputs.clear()
        self.open_button.pack_forget()
        self.convert_button.set_enabled(False)
        self.format_button.set_enabled(False)
        self.add_button.set_enabled(False)
        self.clear_button.set_enabled(False)
        self.browse_button.set_enabled(False)
        self.progress_fraction = 0
        self.draw_progress()
        for path in self.files:
            self.file_states[path] = ("Working", "")
        self.render_files()
        self.status_label.configure(text="Converting files...", fg=self.MUTED)
        threading.Thread(target=self.worker,
                         args=(list(self.files), self.target_var.get(), Path(folder)),
                         daemon=True).start()

    def worker(self, files, target, folder):
        successes, errors = [], []
        for index, path in enumerate(files, 1):
            try:
                outputs = convert(path, target, folder)
                successes.extend(outputs)
                self.events.put(("progress", index, path, "Done", ""))
            except Exception as exc:
                errors.append((path, str(exc)))
                self.events.put(("progress", index, path, "Failed", str(exc)))
        self.events.put(("done", successes, errors, folder))

    def poll(self):
        try:
            while True:
                event = self.events.get_nowait()
                if event[0] == "progress":
                    _, index, path, state, detail = event
                    self.file_states[path] = (state, detail)
                    self.progress_fraction = index / max(1, len(self.files))
                    self.draw_progress()
                    self.render_files()
                    self.status_label.configure(text=f"Processed {index} of {len(self.files)} file(s).",
                                                fg=self.MUTED)
                else:
                    _, successes, errors, folder = event
                    self.running = False
                    self.outputs = successes
                    self.add_button.set_enabled(True)
                    self.browse_button.set_enabled(True)
                    self.refresh_targets()
                    self.render_files()
                    if successes:
                        self.open_button.pack(side="right")
                    if errors:
                        self.status_label.configure(
                            text=f"{len(successes)} output(s) created · {len(errors)} file(s) failed.",
                            fg=self.ERROR)
                    else:
                        self.status_label.configure(
                            text=f"Done. {len(successes)} output file(s) created.",
                            fg=self.SUCCESS)
        except queue.Empty:
            pass
        self.root.after(100, self.poll)

    def open_output(self):
        folder = self.folder_var.get().strip()
        if folder and Path(folder).is_dir():
            os.startfile(folder)



def check_runtime() -> None:
    required = ("PIL", "fitz", "docx", "reportlab", "imageio_ffmpeg")
    missing = []
    for module in required:
        try:
            importlib.import_module(module)
        except ImportError:
            missing.append(module)
    if missing:
        names = ", ".join(missing)
        raise RuntimeError(
            f"Missing RayConvert components: {names}. "
            "Extract the complete RayConvert.zip and keep the vendor folder "
            "beside rayconvert.py."
        )


def main():
    parser = argparse.ArgumentParser(description="RayConvert file converter")
    parser.add_argument("--convert", nargs="+", type=Path, metavar="FILE")
    parser.add_argument("--to", help="Target file extension")
    parser.add_argument("--output-dir", type=Path)
    args = parser.parse_args()
    try:
        check_runtime()
    except RuntimeError as exc:
        if args.convert:
            parser.exit(2, f"{exc}\n")
        root = tk.Tk()
        root.withdraw()
        messagebox.showerror("RayConvert setup", str(exc))
        root.destroy()
        return
    if args.convert:
        if not args.to or not args.output_dir:
            parser.error("--to and --output-dir are required with --convert")
        results = []
        for path in args.convert:
            results.extend(str(p) for p in convert(path, args.to, args.output_dir))
        print(json.dumps(results, ensure_ascii=False))
    else:
        root = tk.Tk()
        RayConvertApp(root)
        root.mainloop()

if __name__ == "__main__":
    main()








